"""Recalculated workbook → data/averra-data.json (the single source of numbers for the deck and the website).

Usage (from repo root):
    python model/build_model.py              # writes model/AVERRA_Financial_Model.xlsx
    python <xlsx skill>/scripts/recalc.py model/AVERRA_Financial_Model.xlsx   # or open & save in Excel
    python scripts/export_data.py
"""
import json, sys
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
wb = load_workbook(ROOT / "model" / "AVERRA_Financial_Model.xlsx", data_only=True)
g = lambda sh, ref: wb[sh][ref].value

def need(v, where):
    if v is None:
        sys.exit(f"Missing value at {where} — recalculate the workbook first.")
    return v

FC = wb["Forecast"]
NM = 36
months = [FC.cell(row=3, column=4 + i).value.replace("\n", " ") for i in range(NM)]
row = lambda r: [round(FC.cell(row=r, column=4 + i).value or 0) for i in range(NM)]
yr = lambda r: {k: round(FC.cell(row=r, column=4 + NM + j).value or 0) for j, k in enumerate(("y1", "p1", "p2", "y2", "y3"))}

ts_cols = {"A": "B", "B": "C", "C": "D"}
TC = "TremoSense Costs"
d = {
    "asOf": "2026-10-12",
    "fx": g("Assumptions", "C42") if False else None,
    "tremosense": {
        "scenarios": {k: {
            "label": {"A": "Prototype / very low volume", "B": "Small batch (100–500)", "C": "Scaled, after verified reductions"}[k],
            "lines": [{"item": wb[TC][f"A{r}"].value, "cost": round(wb[TC][f"{c}{r}"].value)} for r in range(4, 12)],
            "cogs": round(need(wb[TC][f"{c}12"].value, f"{TC}!{c}12")),
            "electronics": round(wb[TC][f"{c}13"].value),
            "price": wb[TC][f"{c}17"].value,
            "exGst": round(wb[TC][f"{c}18"].value),
            "channel": round(wb[TC][f"{c}19"].value),
            "netRevenue": round(wb[TC][f"{c}20"].value),
            "warranty": round(wb[TC][f"{c}22"].value),
            "fulfilment": round(wb[TC][f"{c}23"].value),
            "contribution": round(wb[TC][f"{c}24"].value),
            "contributionMargin": round(wb[TC][f"{c}25"].value, 4),
            "grossMargin": round(wb[TC][f"{c}27"].value, 4),
            "breakEvenUnitsPerMonth": round(wb[TC][f"{c}26"].value, 1),
        } for k, c in ts_cols.items()},
        "prototypePartsCost": round(g(TC, "B14")),
        "volumeTargetElectronics": 2000,
        "priceMatrix": [{"price": wb["Price Matrix"][f"A{r}"].value,
                         "A": round(wb["Price Matrix"][f"B{r}"].value), "Am": round(wb["Price Matrix"][f"C{r}"].value, 4),
                         "B": round(wb["Price Matrix"][f"D{r}"].value), "Bm": round(wb["Price Matrix"][f"E{r}"].value, 4),
                         "C": round(wb["Price Matrix"][f"F{r}"].value), "Cm": round(wb["Price Matrix"][f"G{r}"].value, 4),
                         "verdict": wb["Price Matrix"][f"H{r}"].value} for r in range(4, 11)],
        "channels": [{"channel": wb["Price Matrix"][f"A{r}"].value, "net": round(wb["Price Matrix"][f"B{r}"].value),
                      "contribution": round(wb["Price Matrix"][f"C{r}"].value), "margin": round(wb["Price Matrix"][f"D{r}"].value, 4),
                      "note": wb["Price Matrix"][f"E{r}"].value} for r in range(16, 21)],
        "rental": {"netPerMonth": round(g("Price Matrix", "B22")), "paybackMonths": round(g("Price Matrix", "B23"), 1)},
        "launchPrice": g("Assumptions", "C12"), "targetPrice": g("Assumptions", "C13"), "postCostPrice": g("Assumptions", "C14"),
    },
    "unitEconomics": {},
    "market": {},
    "forecast": {"months": months, "costs": row(32), "cashIn": row(43), "net": row(45), "cumulative": row(46),
                 "cashA": row(48), "cashB": row(49), "cashC": row(50), "cashD": row(51), "tsUnits": row(34), "kits": row(37), "psUnits": row(39), "revTs": row(36), "revKits": row(38), "revPs": row(40), "revSub": row(42),
                 "years": {"costs": yr(32), "cashIn": yr(43), "net": yr(45), "tsUnits": yr(34), "psUnits": yr(39)}},
    "headline": {k: (FC[f"D{r}"].value if isinstance(FC[f"D{r}"].value, str) else round(FC[f"D{r}"].value, 4 if k == "y3Margin" else 0)) for k, r in (
                 ("costs12", 60), ("cashIn12", 61), ("grossBurn12", 62), ("netBurn12", 63), ("phase1Costs", 64), ("phase1Gap", 65),
                 ("fullGap", 66), ("peakGapMonth", 67), ("fullRequirement", 68), ("phase1Requirement", 69), ("phase2Requirement", 70),
                 ("breakEvenMonth", 71), ("y2CashIn", 72), ("y2Net", 73), ("y3CashIn", 74), ("y3Profit", 75), ("y3Margin", 76), ("m36Profit", 77),
                 ("splitShared", 78), ("splitTremoSense", 79), ("splitParkinSense", 80), ("fixedMonthlyY3", 81), ("fixedMonthlyY1", 82))},
    "valuation": {"personMonths": g("Valuation", "B7"), "rate": g("Valuation", "B8"), "materials": g("Valuation", "B9"), "preMoney": g("Valuation", "B10"), "floor": g("Valuation", "B10"),
                  "berkus": g("Valuation", "F9"), "proposed": g("Valuation", "F10"), "berkusFactors": [{"factor": wb["Valuation"][f"E{r}"].value, "value": wb["Valuation"][f"F{r}"].value, "evidence": wb["Valuation"][f"G{r}"].value} for r in range(4, 9)],
                  "options": [{"pre": round(wb["Valuation"][f"A{r}"].value), "investor": round(wb["Valuation"][f"B{r}"].value, 4), "founders": round(wb["Valuation"][f"C{r}"].value, 4), "basis": wb["Valuation"][f"D{r}"].value} for r in range(22, 26)],
                  "workstreams": [{"name": wb["Valuation"][f"A{r}"].value, "pm": wb["Valuation"][f"B{r}"].value} for r in (4, 5, 6)],
                  "targets": [{"equity": round(wb["Valuation"][f"A{r}"].value), "p15": round(wb["Valuation"][f"B{r}"].value),
                               "p20": round(wb["Valuation"][f"C{r}"].value), "dilutionAtP0": round(wb["Valuation"][f"D{r}"].value, 4)} for r in range(14, 18)],
                  },
    "funding": {},
    "useOfFunds": [{"category": wb["Use of Funds"][f"A{r}"].value, "phase1": round(wb["Use of Funds"][f"B{r}"].value), "phase2": round(wb["Use of Funds"][f"C{r}"].value)} for r in range(4, 11)],
    "competitors": [{"name": wb["Competitors"][f"A{r}"].value, "segment": wb["Competitors"][f"B{r}"].value, "price": wb["Competitors"][f"C{r}"].value,
                     "currency": wb["Competitors"][f"D{r}"].value, "inr": round(wb["Competitors"][f"E{r}"].value), "recurring": wb["Competitors"][f"F{r}"].value,
                     "what": wb["Competitors"][f"G{r}"].value, "source": wb["Competitors"][f"H{r}"].value, "checked": wb["Competitors"][f"I{r}"].value} for r in range(4, 11)],
    "sources": [{"label": wb["Sources"][f"A{r}"].value, "url": wb["Sources"][f"B{r}"].value} for r in range(3, wb["Sources"].max_row + 1) if wb["Sources"][f"A{r}"].value],
}
A = {wb["Assumptions"][f"A{r}"].value: wb["Assumptions"][f"C{r}"].value for r in range(4, wb["Assumptions"].max_row + 1)}
d["assumptions"] = A
d["fx"] = A["F6"]
U = wb["Unit Economics"]
for key, c in (("tremosense", "B"), ("parkinsenseDevice", "C"), ("parkinsenseSub", "D"), ("ruoKit", "E")):
    d["unitEconomics"][key] = {k: round(U[f"{c}{r}"].value, 4 if "m" in k.lower() and "argin" in k else 0) for k, r in
                               (("price", 4), ("exGst", 5), ("channel", 6), ("net", 7), ("cogs", 8), ("warranty", 9), ("fulfilment", 10),
                                ("contribution", 11), ("contributionMargin", 12), ("grossMargin", 13))}
for key, c in (("tremosense", "B"), ("parkinsense", "C")):
    d["unitEconomics"][key + "Ltv"] = {"cac": round(U[f"{c}17"].value), "lifetimeMonths": round(U[f"{c}18"].value, 1), "ltv": round(U[f"{c}19"].value),
                                       "ltvCac": round(U[f"{c}20"].value, 2), "paybackMonths": round(U[f"{c}21"].value, 1)}
MK = wb["Market"]
for key, c in (("parkinsense", "B"), ("tremosense", "C")):
    d["market"][key] = {"customers": round(MK[f"{c}7"].value), "revPerCustomer": round(MK[f"{c}8"].value), "tam": round(MK[f"{c}9"].value),
                        "samCustomers": round(MK[f"{c}10"].value), "sam": round(MK[f"{c}11"].value), "somCustomers": round(MK[f"{c}15"].value),
                        "som": round(MK[f"{c}16"].value), "somShare": round(MK[f"{c}17"].value, 4)}
GROUPS = [("People (founders, support, sales, app engineer)", [4, 16, 23, 24, 25]),
          ("Product units: inventory, pilot batch, study units", [10, 15, 19, 20, 27]),
          ("Customer acquisition and clinic outreach", [21, 28]),
          ("Validation and research studies", [11, 13, 14, 22]),
          ("Regulatory, quality and EMC", [12, 17, 18]),
          ("Prototyping, design, tooling and patent", [5, 6, 9, 26]),
          ("Software, cloud, legal and admin", [7, 8, 29]),
          ("Contingency (10%)", [31])]
def ysum(rows):
    out = {k: 0 for k in ("y1", "p1", "p2", "y2", "y3")}
    for r in rows:
        for k, v in yr(r).items(): out[k] += v
    return out
d["forecast"]["costLines"] = [dict(label=l, **ysum(rs)) for l, rs in GROUPS]
d["forecast"]["revLines"] = [dict(label="TremoSense", units=yr(34), **yr(36)), dict(label="ParkinSense research kits", units=yr(37), **yr(38)),
                             dict(label="ParkinSense devices", units=yr(39), **yr(40)), dict(label="ParkinSense subscriptions", units=None, **yr(42))]
d["market"]["general"] = {"customers": round(MK["E7"].value), "revPerCustomer": round(MK["E8"].value), "tam": round(MK["E9"].value), "samCustomers": round(MK["E10"].value), "sam": round(MK["E11"].value), "somCustomers": round(MK["E15"].value), "som": round(MK["E16"].value)}
d["market"]["pd60"] = round(MK["B5"].value); d["market"]["et"] = round(MK["C6"].value)
d["market"]["totals"] = {"tam": round(MK["B19"].value), "sam": round(MK["B20"].value), "som": round(MK["B21"].value)}
FS = wb["Funding Scenarios"]
for k, c in zip("ABCD", "CDEF"):
    d["funding"][k] = {"requirement": round(FS[f"{c}4"].value), "grantStatus": FS[f"{c}5"].value, "grant": round(FS[f"{c}6"].value),
                       "equity": round(FS[f"{c}7"].value), "pre": round(FS[f"{c}8"].value), "post": round(FS[f"{c}9"].value),
                       "investor": round(FS[f"{c}10"].value, 4), "esop": FS[f"{c}11"].value, "founders": round(FS[f"{c}12"].value, 4),
                       "runway": FS[f"{c}13"].value, "milestones": FS[f"{c}14"].value, "risk": FS[f"{c}15"].value,
                       "zeroCash": str(FS[f"{c}16"].value).replace("\n", " ")}
d["funding"]["D"]["tranche2"] = {"equity": round(FS["F20"].value), "pre": round(FS["F21"].value), "investor": round(FS["F22"].value, 4),
                                 "foundersAfterBoth": round(FS["F23"].value, 4), "investorsAfterBoth": round(FS["F24"].value, 4)}
d["funding"]["recommended"] = "C"
out = ROOT / "data" / "averra-data.json"
out.parent.mkdir(exist_ok=True)
out.write_text(json.dumps(d, indent=1, ensure_ascii=False))
print("wrote", out)
