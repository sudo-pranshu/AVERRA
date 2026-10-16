"""Renders the recalculated workbook to public/model.html (values as computed), one tab per sheet."""
import html, datetime
from pathlib import Path
from openpyxl import load_workbook
ROOT = Path(__file__).resolve().parents[1]
wb = load_workbook(ROOT / "model/AVERRA_Financial_Model.xlsx", data_only=True)
def fmt(c):
    v = c.value
    if v is None: return ""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        f = c.number_format or ""
        if "%" in f: return f"{v*100:.1f}%"
        if abs(v) >= 1000 or float(v).is_integer(): return f"{v:,.0f}"
        return f"{v:,.2f}"
    if isinstance(v, (datetime.date, datetime.datetime)): return v.strftime("%d %b %Y")
    return str(v)
tabs, panes = [], []
for i, ws in enumerate(wb.worksheets):
    rows = []
    for r in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=min(ws.max_column, 45)):
        cells = [fmt(c) for c in r]
        if not any(cells): rows.append("<tr><td>&nbsp;</td></tr>"); continue
        tds = "".join(f'<td class="{"n" if (c.value is not None and isinstance(c.value,(int,float))) else ""}{" b" if c.font and c.font.b else ""}">{html.escape(t)}</td>' for c, t in zip(r, cells))
        rows.append(f"<tr>{tds}</tr>")
    tabs.append(f'<button onclick="show({i})" id="t{i}">{html.escape(ws.title)}</button>')
    panes.append(f'<div class="pane" id="p{i}"><table>{"".join(rows)}</table></div>')
doc = f"""<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>AVERRA financial model</title>
<style>body{{margin:0;font:13px/1.4 -apple-system,Segoe UI,Arial,sans-serif;background:#fff;color:#0f172a}}
nav{{position:sticky;top:0;background:#f1f5f9;border-bottom:1px solid #cbd5e1;padding:6px;display:flex;flex-wrap:wrap;gap:4px;z-index:2}}
nav button{{border:1px solid #cbd5e1;background:#fff;border-radius:6px;padding:4px 10px;cursor:pointer;font:inherit}}
nav button.on{{background:#0e7490;color:#fff;border-color:#0e7490}}
.pane{{display:none;overflow:auto;padding:8px}} table{{border-collapse:collapse}}
td{{border:1px solid #e2e8f0;padding:3px 7px;white-space:nowrap;max-width:520px;overflow:hidden;text-overflow:ellipsis}} td.n{{text-align:right;font-variant-numeric:tabular-nums}} td.b{{font-weight:600}}</style></head>
<body><nav>{"".join(tabs)}</nav>{"".join(panes)}
<script>function show(i){{document.querySelectorAll('.pane').forEach((p,k)=>p.style.display=k==i?'block':'none');document.querySelectorAll('nav button').forEach((b,k)=>b.className=k==i?'on':'')}}show(0)</script></body></html>"""
(ROOT / "public/model.html").write_text(doc)
print("model.html", len(doc), "bytes,", len(wb.worksheets), "sheets")
